"""
Extracts BOQ (Bill of Quantities) line items from an uploaded Excel, PDF, or
Word file.

BOQs are notoriously inconsistent in layout, so this uses a header-detection
heuristic rather than assuming fixed columns: it scans rows for one that
looks like a table header (contains recognisable column names such as
"description", "qty", "unit"), then reads every following row as a line item
using whichever of those columns it found. Rows that don't have a
description are skipped (title rows, section headings, blank spacing rows).

This is intentionally conservative — it will not catch every BOQ layout.
Anything it gets wrong or misses is easy to fix by hand afterwards in the
line items table, and any parsing issues are returned as `warnings` so the
user knows to double check.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import List, Optional, Tuple

ITEM_NO_KEYWORDS = ("item no", "item", "no.", "ref", "pos")
DESCRIPTION_KEYWORDS = ("description", "particulars", "specification", "item description")
UNIT_KEYWORDS = ("unit", "uom", "unit of measure")
QUANTITY_KEYWORDS = ("qty", "quantity", "no. off", "number")


@dataclass
class ParsedLineItem:
    item_no: Optional[str]
    description: str
    unit: Optional[str]
    quantity: Optional[float]


def _normalize_cell(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _find_header_row(rows: List[List[str]]) -> Optional[Tuple[int, dict]]:
    for i, row in enumerate(rows):
        normalized = [_normalize_cell(c) for c in row]
        col_map = {}
        for idx, cell in enumerate(normalized):
            if not cell:
                continue
            if "description" in col_map and "unit" in col_map and "quantity" in col_map:
                break
            if any(k in cell for k in DESCRIPTION_KEYWORDS) and "description" not in col_map:
                col_map["description"] = idx
            elif any(k == cell or cell.startswith(k) for k in UNIT_KEYWORDS) and "unit" not in col_map:
                col_map["unit"] = idx
            elif any(k in cell for k in QUANTITY_KEYWORDS) and "quantity" not in col_map:
                col_map["quantity"] = idx
            elif any(cell == k or cell.startswith(k) for k in ITEM_NO_KEYWORDS) and "item_no" not in col_map:
                col_map["item_no"] = idx
        if "description" in col_map:
            return i, col_map
    return None


def _try_float(value) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _rows_to_line_items(rows: List[List[str]]) -> Tuple[List[ParsedLineItem], List[str]]:
    warnings: List[str] = []
    header = _find_header_row(rows)
    if header is None:
        warnings.append(
            "Could not find a header row with a 'Description' column - no line items were extracted "
            "from this sheet/page. You can still add line items manually."
        )
        return [], warnings

    header_idx, col_map = header
    items: List[ParsedLineItem] = []
    for row in rows[header_idx + 1 :]:
        if col_map["description"] >= len(row):
            continue
        description = row[col_map["description"]]
        if description is None or str(description).strip() == "":
            continue
        description = str(description).strip()
        if description.lower() in ("description", "total", "subtotal", "grand total"):
            continue

        unit_idx = col_map.get("unit")
        qty_idx = col_map.get("quantity")
        item_no_idx = col_map.get("item_no")

        items.append(
            ParsedLineItem(
                item_no=str(row[item_no_idx]).strip() if item_no_idx is not None and item_no_idx < len(row) and row[item_no_idx] not in (None, "") else None,
                description=description,
                unit=str(row[unit_idx]).strip() if unit_idx is not None and unit_idx < len(row) and row[unit_idx] not in (None, "") else None,
                quantity=_try_float(row[qty_idx]) if qty_idx is not None and qty_idx < len(row) else None,
            )
        )

    if not items:
        warnings.append("A header row was found but no data rows followed it.")

    return items, warnings


def parse_excel(content: bytes) -> Tuple[List[ParsedLineItem], List[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    all_items: List[ParsedLineItem] = []
    all_warnings: List[str] = []
    for sheet in wb.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        if not rows:
            continue
        items, warnings = _rows_to_line_items(rows)
        all_items.extend(items)
        all_warnings.extend(f"[{sheet.title}] {w}" for w in warnings)
    return all_items, all_warnings


def parse_pdf(content: bytes) -> Tuple[List[ParsedLineItem], List[str]]:
    import pdfplumber

    all_items: List[ParsedLineItem] = []
    all_warnings: List[str] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table in tables:
                items, warnings = _rows_to_line_items(table)
                all_items.extend(items)
                all_warnings.extend(f"[page {page_no}] {w}" for w in warnings)
    return all_items, all_warnings


def parse_docx(content: bytes) -> Tuple[List[ParsedLineItem], List[str]]:
    import docx

    document = docx.Document(BytesIO(content))
    all_items: List[ParsedLineItem] = []
    all_warnings: List[str] = []
    for t_idx, table in enumerate(document.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        items, warnings = _rows_to_line_items(rows)
        all_items.extend(items)
        all_warnings.extend(f"[table {t_idx + 1}] {w}" for w in warnings)
    return all_items, all_warnings


def parse_boq(filename: str, content: bytes) -> Tuple[List[ParsedLineItem], List[str]]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return parse_excel(content)
    if lower.endswith(".xls"):
        raise ValueError("Legacy .xls files aren't supported - please re-save the BOQ as .xlsx and re-upload.")
    if lower.endswith(".pdf"):
        return parse_pdf(content)
    if lower.endswith(".docx"):
        return parse_docx(content)
    raise ValueError(f"Unsupported BOQ file type: {filename}. Use .xlsx, .pdf, or .docx.")
